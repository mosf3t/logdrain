from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelReportGenerator:
    """Generates a professional diagnostic Excel workbook for Universal Robots Polyscope logs."""

    def __init__(self, records: List[Any], clusters: List[Dict[str, Any]]):
        self.records = records
        self.clusters = clusters

    def generate_workbook(self, output_path: str) -> None:
        wb = Workbook()
        
        # Setup styles
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        border_thin = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        
        # -------------------------------------------------------------
        # Sheet 1: Executive Summary / Clusters
        # -------------------------------------------------------------
        ws_summary = wb.active
        ws_summary.title = "Cluster Summary"
        ws_summary.views.sheetView[0].showGridLines = True

        summary_headers = ["Cluster ID", "Template Text", "Occurrence Count"]
        ws_summary.append(summary_headers)

        for col_num in range(1, len(summary_headers) + 1):
            cell = ws_summary.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for cluster in self.clusters:
            row = [cluster["cluster_id"], cluster["template"], cluster["size"]]
            ws_summary.append(row)
            r_idx = ws_summary.max_row
            for c_idx in range(1, 4):
                c = ws_summary.cell(row=r_idx, column=c_idx)
                c.border = border_thin
                if c_idx == 2:
                    c.alignment = Alignment(wrap_text=True, vertical="top")

        # -------------------------------------------------------------
        # Sheet 2: Parsed Log Records
        # -------------------------------------------------------------
        ws_logs = wb.create_sheet(title="Parsed Logs")
        ws_logs.views.sheetView[0].showGridLines = True

        log_headers = ["ID", "Timestamp", "Severity", "Thread", "Logger Class", "Message Body", "Malformed"]
        ws_logs.append(log_headers)

        for col_num in range(1, len(log_headers) + 1):
            cell = ws_logs.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for rec in self.records:
            row = [
                rec.record_id,
                str(rec.timestamp) if rec.timestamp else "N/A",
                rec.severity,
                rec.thread,
                rec.logger_class,
                rec.message_body,
                str(rec.is_malformed)
            ]
            ws_logs.append(row)
            r_idx = ws_logs.max_row
            for c_idx in range(1, 8):
                c = ws_logs.cell(row=r_idx, column=c_idx)
                c.border = border_thin
                # Enable text wrapping for message body (column 6) and logger class (column 5)
                if c_idx in [5, 6]:
                    c.alignment = Alignment(wrap_text=True, vertical="top")

        # Apply controlled column widths (capped so massive scripts don't stretch layout)
        for sheet in [ws_summary, ws_logs]:
            for col in sheet.columns:
                max_len = 0
                for cell in col:
                    val_str = str(cell.value or '')
                    # If it's a multi-line body, consider the longest line rather than total string length
                    lines = val_str.splitlines()
                    line_max = max((len(l) for l in lines), default=0)
                    if line_max > max_len:
                        max_len = line_max
                
                col_letter = get_column_letter(col[0].column)
                # Cap width between 12 and 60 to keep things readable
                sheet.column_dimensions[col_letter].width = max(min(max_len + 3, 60), 12)

        wb.save(output_path)
